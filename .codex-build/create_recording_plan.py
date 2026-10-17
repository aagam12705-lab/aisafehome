from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path('/Users/aagamjain/Desktop/ai-safehome/output/documents/SafeHome_AI_Recording_Plan_and_Script.docx')
OUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = '203247'
TEAL = '1E5D78'
PALE = 'E3EEF1'
GRID = 'D9D9D9'
ALT = 'F5F7F8'
GRAY = RGBColor(89, 103, 122)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def borders(table, color=GRID):
    tblPr = table._tbl.tblPr
    el = tblPr.find(qn('w:tblBorders'))
    if el is None:
        el = OxmlElement('w:tblBorders')
        tblPr.append(el)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = OxmlElement(f'w:{edge}')
        tag.set(qn('w:val'), 'single')
        tag.set(qn('w:sz'), '4')
        tag.set(qn('w:color'), color)
        el.append(tag)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'), 'dxa')

def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)

def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f'Heading {level}')
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    return p

def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(text)
    p.paragraph_format.space_after = Pt(3)
    return p

def add_number(doc, text):
    p = doc.add_paragraph(style='List Number')
    p.add_run(text)
    p.paragraph_format.space_after = Pt(3)
    return p

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5); sec.page_height = Inches(11)
sec.top_margin = Inches(0.65); sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.72); sec.right_margin = Inches(0.72)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(10.5); styles['Normal'].font.color.rgb = RGBColor(44,62,80)
styles['Normal'].paragraph_format.space_after = Pt(6)
styles['Title'].font.name = 'Aptos Display'; styles['Title'].font.size = Pt(28); styles['Title'].font.bold = True; styles['Title'].font.color.rgb = RGBColor(0,0,0)
styles['Subtitle'].font.name = 'Aptos'; styles['Subtitle'].font.size = Pt(13); styles['Subtitle'].font.color.rgb = GRAY
for name, size in [('Heading 1',18),('Heading 2',13)]:
    styles[name].font.name = 'Aptos Display'; styles[name].font.size = Pt(size); styles[name].font.bold = True; styles[name].font.color.rgb = RGBColor(0,0,0)

p = doc.add_paragraph(style='Title'); p.add_run('SafeHome AI Recording Plan and Script')
p = doc.add_paragraph(style='Subtitle'); p.add_run('Three minute Congressional App Challenge video')
p = doc.add_paragraph()
p.add_run('Purpose  ').bold = True
p.add_run('This plan fits the complete Congressional App Challenge story into three minutes. It identifies the creator, purpose, target audience, originality, tools, technical challenge, validation, working features, accessibility, and limits. Six short slides provide context while four app clips prove the experience works.')

add_heading(doc, 'Final video structure', 1)
table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [Inches(1.05), Inches(1.05), Inches(1.45), Inches(3.55)]
hdr = table.rows[0].cells
for i, text in enumerate(['Time','Media','Length','Purpose']):
    hdr[i].width = widths[i]; hdr[i].text = text; shade(hdr[i], TEAL); set_cell_margins(hdr[i])
    for r in hdr[i].paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255)
set_repeat_header(table.rows[0])
rows = [
('0:00-0:12','Slide 1','12 sec','Open with one older adult, one familiar room, and one overlooked hazard.'),
('0:12-0:32','Clip A','20 sec','Start a Quick Safety Check and upload a photo.'),
('0:32-0:44','Slide 2','12 sec','Show the fast path and the optional saved path.'),
('0:44-1:14','Clip B','30 sec','Connect marked evidence to questions and the score.'),
('1:14-1:30','Slide 3','16 sec','Explain how the interface, AI, rules, database, and email fit together.'),
('1:30-1:50','Clip C','20 sec','Show the action plan and evidence of improvement.'),
('1:50-2:05','Slide 4','15 sec','Explain the bounded scoring system.'),
('2:05-2:20','Slide 5','15 sec','Present the testing result and its limits.'),
('2:20-2:45','Clip D','25 sec','Show saved progress, sharing, and accessibility.'),
('2:45-3:00','Slide 6','15 sec','Return to the room and close on a clear next step.'),
]
for ridx, row in enumerate(rows, start=1):
    cells = table.add_row().cells
    for i, text in enumerate(row):
        cells[i].width = widths[i]; cells[i].text = text; cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cells[i], ALT if ridx % 2 == 0 else 'FFFFFF'); set_cell_margins(cells[i])
        if i < 3: cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
borders(table)

doc.add_page_break()
add_heading(doc, 'Exact clip recording instructions', 1)

add_heading(doc, 'Preparation before every clip', 2)
add_bullet(doc, 'Record at 1920 by 1080 and 60 frames per second with the browser at 100 percent zoom.')
add_bullet(doc, 'Use light mode, the same staged bedroom photo, and the same browser size throughout.')
add_bullet(doc, 'Hide bookmarks, notifications, unrelated tabs, and personal email addresses.')
add_bullet(doc, 'Record the app silently and add narration afterward.')
add_bullet(doc, 'Leave one second still at the beginning and end of every clip.')
add_bullet(doc, 'Record each clip twice and retain the smoother take.')
add_bullet(doc, 'Treat each lettered clip as a short sequence of prepared shots. Stop recording, open the next page, and begin a new shot instead of racing through navigation.')
add_bullet(doc, 'Remove loading time, file-picker time, typing, sign-in, and long scrolls during editing. Do not speed up the footage.')

add_heading(doc, 'Clip A  0 12 to 0 32  20 seconds', 2)
add_number(doc, 'Shot A1, about eight seconds: begin on the landing page, click Quick Safety Check, choose Bedroom, and click Continue at a normal pace.')
add_number(doc, 'Stop recording. Open the upload page and load the staged photo off camera.')
add_number(doc, 'Shot A2, about twelve seconds: begin with the photo already visible, pause on the quality result, click Analyze Photo, and cut as analysis begins.')
p = doc.add_paragraph(); p.add_run('Narration  ').bold = True; p.add_run('I wanted the first check to feel simple, so no account is required. Choose the room, upload one photo, and the app checks orientation, lighting, contrast, and blur before analysis.')

add_heading(doc, 'Clip B  0 44 to 1 14  30 seconds', 2)
add_number(doc, 'Shot B1, about twelve seconds: begin on the AI results page, hold the marked image, then make one slow scroll to the first hazard card.')
add_number(doc, 'Stop recording. Open the follow up questions page off camera.')
add_number(doc, 'Shot B2, about seven seconds: show one question, choose an answer, and cut immediately after the selection registers.')
add_number(doc, 'Stop recording. Open the completed score page off camera.')
add_number(doc, 'Shot B3, about eleven seconds: hold on the numerical score and colored risk bar, then make one small scroll to the first priority action if it is directly below.')
p = doc.add_paragraph(); p.add_run('Narration  ').bold = True; p.add_run('I designed blue markers to show exactly what the AI noticed instead of giving a vague warning. The app asks only about details the camera cannot prove, then combines those answers with validated hazards to calculate the score and three priority actions.')

add_heading(doc, 'Clip C  1 30 to 1 50  20 seconds', 2)
add_number(doc, 'Shot C1, about ten seconds: begin on the safety plan and make one slow scroll that reveals the three priority actions.')
add_number(doc, 'Stop recording. Complete the recheck off camera and open the finished before and after comparison.')
add_number(doc, 'Shot C2, about ten seconds: show the score change and pause on Resolved, Still Present, and New. End with the comparison still on screen.')
p = doc.add_paragraph(); p.add_run('Narration  ').bold = True; p.add_run('I did not want the app to stop at detection. Each score leads to three practical actions. After a fix, another photo shows the score change, what was resolved, what remains, and any new concern.')

add_heading(doc, 'Clip D  2 20 to 2 45  25 seconds', 2)
add_number(doc, 'Prepare a signed in demo account with saved checks before recording.')
add_number(doc, 'Shot D1, about ten seconds: begin on Room by Room Stats, hold the unique rooms count, then make one slow scroll to the score trend.')
add_number(doc, 'Stop recording. Open an Email this summary section off camera.')
add_number(doc, 'Shot D2, about six seconds: click Add new email address and show the new empty recipient field. Do not type or send anything.')
add_number(doc, 'Stop recording. Return to a clean page and open Accessibility off camera.')
add_number(doc, 'Shot D3, about nine seconds: change one visible setting, then toggle Read Aloud or Reduced Motion. End with the panel visible.')
p = doc.add_paragraph(); p.add_run('Narration  ').bold = True; p.add_run('For families who want to keep improving, signed-in checks track each room over time. Summaries can reach five recipients. Larger text, four color modes, Read Aloud, reduced motion, and a mobile layout keep the experience senior-friendly.')

add_heading(doc, 'Complete narration and screen directions', 1)

segments = [
('0:00-0:12  Slide 1',
 'A fall can begin with one overlooked cord. I am Aagam Jain, and I built SafeHome AI to help older adults and caregivers notice everyday risks sooner.',
 'Hold the opening slide. Your name, the audience, and the one-sentence purpose must remain readable.'),
('0:12-0:32  Clip A',
 'I wanted the first check to feel simple, so no account is required. Choose the room, upload one photo, and the app checks orientation, lighting, contrast, and blur before analysis.',
 'Use two shots. First, move from the landing page through Bedroom selection. Cut. Start again with the uploaded photo already visible, show its quality result, and click Analyze.'),
('0:32-0:44  Slide 2',
 'Instead of a long checklist, SafeHome AI starts with the room and asks only what the image cannot answer. An account adds progress, comparisons, and sharing.',
 'Hold the slide long enough for both paths to be read. Use a simple dissolve from the previous clip.'),
('0:44-1:14  Clip B',
 'I designed blue markers to show exactly what the AI noticed instead of giving a vague warning. The app asks only about details the camera cannot prove, then combines those answers with validated hazards to calculate the score and three priority actions.',
 'Use three prepared shots: marked evidence and one hazard card, one follow-up question, then the completed score and risk bar. Cut between pages instead of navigating on camera.'),
('1:14-1:30  Slide 3',
 'I built the app with Python, Streamlit, and CSS. Vision AI locates evidence, while my validation layer checks categories, clamps points, and uses fallbacks. Supabase stores history, Brevo sends text-only summaries, and photos are never saved.',
 'Follow the six top-row steps from left to right. Then point to the Supabase, Brevo, and active-session branches before ending on the failure-handling line.'),
('1:30-1:50  Clip C',
 'I did not want the app to stop at detection. Each score leads to three practical actions. After a fix, another photo shows the score change, what was resolved, what remains, and any new concern.',
 'Use two prepared shots: the three recommended actions, then the completed comparison. Perform the photo upload and recheck off camera.'),
('1:50-2:05  Slide 4',
 'The hardest problem was inconsistent AI scoring. I solved it by limiting hazards to one through fifteen points, adding confirmed follow-up concerns, using category fallbacks, and capping the room at one hundred.',
 'Hold the score slide. The three inputs and low, moderate, and high ranges must remain readable.'),
('2:05-2:20  Slide 5',
 'For validation, I labeled thirty-five hazards across ten staged photos before running the app. It found thirty-four. AI can still miss small details, so every result keeps human review explicit.',
 'Show the full testing table without zooming or panning. Keep the total row and the warning that AI can miss small details visible.'),
('2:20-2:45  Clip D',
 'For families who want to keep improving, signed-in checks track each room over time. Summaries can reach five recipients. Larger text, four color modes, Read Aloud, reduced motion, and a mobile layout keep the experience senior-friendly.',
 'Use three prepared shots: dashboard and trend, Add new email address, then Accessibility. Cut between them. Do not type an address or navigate across the app while recording.'),
('2:45-3:00  Slide 6',
 'I have also used SafeHome AI in more than two hundred rooms belonging to friends and family, and they found the results extremely useful. My goal is to help more families catch risks sooner.',
 'Briefly point to the design decisions and human-review statement, then hold on the QR code. Leave one second of silence before ending.'),
]

for title, narration, direction in segments:
    add_heading(doc, title, 2)
    p = doc.add_paragraph(); p.add_run('Narration  ').bold = True; p.add_run(narration)
    p = doc.add_paragraph(); p.add_run('Screen direction  ').bold = True; p.add_run(direction)

doc.add_page_break()
add_heading(doc, 'Clip recording checklist', 1)
add_bullet(doc, 'Record at 1920 by 1080 and 60 frames per second.')
add_bullet(doc, 'Use the same browser size, zoom level, light theme, staged room, and demo account throughout.')
add_bullet(doc, 'Hide bookmarks, notifications, personal email addresses, and unrelated browser tabs.')
add_bullet(doc, 'Leave about one second still at the beginning and end of every clip for clean edits.')
add_bullet(doc, 'Move the cursor slowly and perform one clear action at a time.')
add_bullet(doc, 'Do not speed up scrolling enough to make labels unreadable.')
add_bullet(doc, 'Pause on the marked image, the final score, the testing result, and the before-and-after comparison.')
add_bullet(doc, 'Record each clip at least twice and keep the smoother take.')
add_bullet(doc, 'Use a staged room photo and demo recipient addresses in the final recording.')

add_heading(doc, 'Editing checklist', 1)
add_bullet(doc, 'Use only Slides 1 through 6 from the presentation. The production-reference slides do not belong in the video.')
add_bullet(doc, 'Keep the final runtime at or below three minutes.')
add_bullet(doc, 'Use simple cuts or short dissolves. Avoid animated templates and decorative transitions.')
add_bullet(doc, 'Keep app clips at normal speed unless a very small trim is needed.')
add_bullet(doc, 'Match each narration block to its listed slide or clip before adding music or captions.')
add_bullet(doc, 'Add readable captions and check that they do not cover buttons, scores, or testing totals.')
add_bullet(doc, 'Keep background music quiet enough that every word remains clear.')
add_bullet(doc, 'Watch the complete export on a phone and a laptop before submitting.')

add_heading(doc, 'Final quality check', 1)
p = doc.add_paragraph()
p.add_run('The finished video should feel like a guided app demonstration, not a slideshow. ').bold = True
p.add_run('The viewer should see the complete path from starting a check to receiving evidence, answering follow-up questions, viewing the score, acting on recommendations, tracking progress, and sharing a summary.')

doc.core_properties.title = 'SafeHome AI Recording Plan and Script'
doc.core_properties.subject = 'Three minute Congressional App Challenge video production plan'
doc.core_properties.author = 'SafeHome AI project team'
doc.save(OUT)
print(OUT)
